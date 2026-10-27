# CONVERSOR DE PNG PARA DOCX (WORD)
# Salva com data e hora para não dar conflito

from docx import Document
from docx.shared import Inches
from PIL import Image
from tkinter import Tk, filedialog
from datetime import datetime
import os

# 1. Pasta Imagens do usuário
pasta_imagens = os.path.join(os.path.expanduser("~"), "Pictures")
print("Pasta de destino:", pasta_imagens)

# 2. Escolhe a imagem
print("Escolha a imagem PNG na janelinha que abriu...")

root = Tk()
root.withdraw()
caminho_png = filedialog.askopenfilename(
    title="Escolha a imagem PNG",
    filetypes=[("Imagens", "*.png *.jpg *.jpeg")]
)

if not caminho_png:
    print("Nenhum arquivo selecionado. Encerrando.")
else:
    print("Arquivo escolhido:", caminho_png)

    # 3. Abre a imagem
    imagem = Image.open(caminho_png)
    largura_px, altura_px = imagem.size

    # 4. Cria o Word
    doc = Document()
    doc.add_heading("Imagem Convertida", level=1)

    largura_pol = min(largura_px / 96, 6)
    doc.add_picture(caminho_png, width=Inches(largura_pol))

    # 5. Nome único com data e hora
    agora = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    nome_docx = os.path.join(pasta_imagens, f"imagem_convertida_{agora}.docx")
    doc.save(nome_docx)

    print("Documento Word salvo em:", nome_docx)