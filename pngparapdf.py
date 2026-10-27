# CONVERSOR DE PNG PARA DOCX - JEITO FÁCIL
# Abre uma janelinha para escolher a imagem

from docx import Document
from docx.shared import Inches
from PIL import Image
from tkinter import Tk, filedialog
import os

# 1. Abre a janelinha para escolher o arquivo
print("Escolha a imagem PNG na janelinha que abriu...")

root = Tk()
root.withdraw()  # esconde a janela principal
caminho_png = filedialog.askopenfilename(
    title="Escolha a imagem PNG",
    filetypes=[("Imagens", "*.png *.jpg *.jpeg")]
)

# 2. Verifica se o usuário escolheu algo
if not caminho_png:
    print("Nenhum arquivo selecionado. Encerrando.")
else:
    print("Arquivo escolhido:", caminho_png)

    # 3. Abre a imagem
    imagem = Image.open(caminho_png)
    largura_px, altura_px = imagem.size

    # 4. Cria o documento Word
    doc = Document()
    doc.add_heading("Imagem Convertida", level=1)

    # 5. Ajusta o tamanho para caber na página
    largura_pol = min(largura_px / 96, 6)

    # 6. Insere a imagem
    doc.add_picture(caminho_png, width=Inches(largura_pol))

    # 7. Salva o DOCX na mesma pasta da imagem
    pasta = os.path.dirname(caminho_png)
    nome_docx = os.path.join(pasta, "resultado.docx")
    doc.save(nome_docx)

    print("Documento criado com sucesso:", nome_docx)